"""
Export Service - Serviço genérico para exportação de dados.

Implementa o padrão de template para exportação em diferentes formatos
(Excel, PDF) mantendo uma interface consistente.
"""
from __future__ import annotations

import io
from abc import ABC
from abc import abstractmethod
from dataclasses import dataclass
from dataclasses import field
from datetime import datetime
from decimal import Decimal
from typing import TYPE_CHECKING
from typing import Any

from django.http import HttpResponse
from django.template.loader import render_to_string

if TYPE_CHECKING:
    from django.db.models import QuerySet


@dataclass
class ExportColumn:
    """Define uma coluna para exportação."""
    field: str  # Nome do campo ou callable
    header: str  # Título da coluna
    width: int = 15  # Largura (para Excel)
    format_func: callable | None = None  # Função para formatar valor
    align: str = "left"  # left, center, right


@dataclass
class ExportConfig:
    """Configuração para exportação."""
    filename: str
    title: str
    columns: list[ExportColumn]
    subtitle: str = ""
    company_name: str = ""
    logo_path: str | None = None
    generated_by: str = ""
    page_orientation: str = "portrait"  # portrait, landscape
    extra_data: dict = field(default_factory=dict)


class BaseExportService(ABC):
    """
    Classe base abstrata para serviços de exportação.

    Fornece a estrutura comum para exportar dados em diferentes formatos.
    Subclasses devem implementar os métodos abstratos para cada formato.
    """

    def __init__(self, config: ExportConfig):
        self.config = config

    @abstractmethod
    def export(self, queryset: "QuerySet") -> HttpResponse:
        """Exporta os dados e retorna uma HttpResponse."""
        pass

    def _get_row_data(self, obj: Any) -> list:
        """Extrai dados de uma linha baseado nas colunas configuradas."""
        row = []
        for col in self.config.columns:
            value = self._get_field_value(obj, col.field)
            if col.format_func:
                value = col.format_func(value)
            row.append(value)
        return row

    def _get_field_value(self, obj: Any, field: str) -> Any:
        """
        Obtém o valor de um campo do objeto.

        Suporta:
        - Campos diretos: 'nome'
        - Campos aninhados: 'categoria__nome'
        - Métodos: 'get_display_name'
        - Callables: lambda obj: obj.calcular_total()
        """
        if callable(field):
            return field(obj)

        # Campos aninhados com double underscore
        parts = field.split("__")
        value = obj

        for part in parts:
            if hasattr(value, part):
                attr = getattr(value, part)
                value = attr() if callable(attr) else attr
            else:
                value = None
                break

        return value

    def _format_value(self, value: Any) -> str:
        """Formata um valor para exibição."""
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%d/%m/%Y %H:%M")
        if isinstance(value, Decimal):
            return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        if isinstance(value, bool):
            return "Sim" if value else "Não"
        return str(value)


class ExcelExportService(BaseExportService):
    """
    Serviço de exportação para Excel (.xlsx).

    Utiliza openpyxl para gerar arquivos Excel formatados
    com suporte a estilos, largura de colunas e cabeçalhos.
    """

    def export(self, queryset: "QuerySet") -> HttpResponse:
        """Gera arquivo Excel e retorna como HttpResponse."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment
            from openpyxl.styles import Border
            from openpyxl.styles import Font
            from openpyxl.styles import PatternFill
            from openpyxl.styles import Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError(
                "openpyxl é necessário para exportação Excel. "
                "Instale com: pip install openpyxl"
            )

        wb = Workbook()
        ws = wb.active
        ws.title = self.config.title[:31]  # Excel limita a 31 caracteres

        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Título do relatório
        row_num = 1
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(self.config.columns))
        title_cell = ws.cell(row=1, column=1, value=self.config.title)
        title_cell.font = Font(bold=True, size=14)
        title_cell.alignment = Alignment(horizontal="center")
        row_num += 1

        # Subtítulo/empresa
        if self.config.company_name or self.config.subtitle:
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(self.config.columns))
            subtitle = f"{self.config.company_name} - {self.config.subtitle}" if self.config.company_name else self.config.subtitle
            ws.cell(row=row_num, column=1, value=subtitle).alignment = Alignment(horizontal="center")
            row_num += 1

        # Data de geração
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(self.config.columns))
        ws.cell(row=row_num, column=1, value=f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
        row_num += 2

        # Cabeçalhos
        for col_idx, col in enumerate(self.config.columns, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=col.header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col_idx)].width = col.width

        row_num += 1

        # Dados
        for obj in queryset:
            row_data = self._get_row_data(obj)
            for col_idx, (col, value) in enumerate(zip(self.config.columns, row_data), 1):
                cell = ws.cell(row=row_num, column=col_idx, value=self._format_value(value))
                cell.border = thin_border
                cell.alignment = Alignment(horizontal=col.align)
            row_num += 1

        # Preparar resposta HTTP
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"{self.config.filename}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class PDFExportService(BaseExportService):
    """
    Serviço de exportação para PDF.

    Utiliza Django templates e weasyprint para gerar PDFs
    formatados com suporte à identidade visual da empresa.
    """

    def __init__(
        self,
        config: ExportConfig,
        template_name: str = "core/exports/pdf_report.html",
    ):
        super().__init__(config)
        self.template_name = template_name

    def export(self, queryset: "QuerySet") -> HttpResponse:
        """Gera arquivo PDF e retorna como HttpResponse."""
        try:
            from weasyprint import HTML
        except ImportError:
            raise ImportError(
                "weasyprint é necessário para exportação PDF. "
                "Instale com: pip install weasyprint"
            )

        # Preparar dados para o template
        rows = []
        for obj in queryset:
            row_data = self._get_row_data(obj)
            # Include alignment info with each cell for template access
            row_cells = [
                {"value": self._format_value(v), "align": col.align}
                for v, col in zip(row_data, self.config.columns)
            ]
            rows.append(row_cells)

        context = {
            "config": self.config,
            "columns": self.config.columns,
            "rows": rows,
            "generated_at": datetime.now(),
            "total_records": queryset.count(),
            **self.config.extra_data,
        }

        # Renderizar HTML
        html_content = render_to_string(self.template_name, context)

        # Converter para PDF
        pdf_file = HTML(string=html_content).write_pdf()

        filename = f"{self.config.filename}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        response = HttpResponse(pdf_file, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    def preview_html(self, queryset: "QuerySet") -> str:
        """Retorna o HTML do relatório para preview."""
        rows = []
        for obj in queryset:
            row_data = self._get_row_data(obj)
            # Include alignment info with each cell for template access
            row_cells = [
                {"value": self._format_value(v), "align": col.align}
                for v, col in zip(row_data, self.config.columns)
            ]
            rows.append(row_cells)

        context = {
            "config": self.config,
            "columns": self.config.columns,
            "rows": rows,
            "generated_at": datetime.now(),
            "total_records": queryset.count(),
            **self.config.extra_data,
        }

        return render_to_string(self.template_name, context)
